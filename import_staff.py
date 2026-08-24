import openpyxl
from app import create_app, db
from app.models import User, Branch
from app.utils.timezone import now_pkt
import re

def import_staff():
    app = create_app()
    with app.app_context():
        wb = openpyxl.load_workbook('CDC_lab.xlsx')
        ws = wb['cdc-staff']
        
        print(f"Reading spreadsheet: {ws.max_row - 1} entries found.")
        
        # We need a starting counter for sequence of employee IDs
        max_emp_seq = 3000
        # Check existing user employee IDs to avoid duplicate employee_id
        for u in User.query.all():
            match = re.match(r'EMP-(\d+)', u.employee_id or '')
            if match:
                seq = int(match.group(1))
                if seq > max_emp_seq:
                    max_emp_seq = seq
        
        user_count = 0
        branch_count = 0
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            email = str(row[1] or '').strip()
            name = str(row[2] or '').strip()
            designation = str(row[3] or '').strip()
            branch_name = str(row[4] or '').strip()
            role_desc = str(row[10] or '').strip() # Current role column
            
            if not name or not email:
                continue
                
            # Generate username from email
            username = email.split('@')[0].lower().replace('.', '_').replace('-', '_')
            
            # Map designation and current role to system role
            # ADMIN, MANAGER, SUPERVISOR, BRANCH_STAFF, RIDER, LAB_STAFF, VERIFIER, RECEPTION_STAFF
            role = 'LAB_STAFF' # Default
            desig_lower = designation.lower()
            role_desc_lower = role_desc.lower()
            
            if 'manager' in desig_lower or 'manager' in role_desc_lower:
                role = 'MANAGER'
            elif 'supervisor' in desig_lower or 'supervisor' in role_desc_lower:
                role = 'SUPERVISOR'
            elif 'reception' in desig_lower or 'reception' in role_desc_lower:
                role = 'BRANCH_STAFF'
            elif 'phlebotomist' in desig_lower or 'phlebotomist' in role_desc_lower:
                role = 'LAB_STAFF'
            elif 'technologist' in desig_lower or 'verifier' in desig_lower:
                role = 'VERIFIER' if 'verifier' in desig_lower or 'verifier' in role_desc_lower else 'LAB_STAFF'
            
            # Clean Branch Name and Code
            if not branch_name:
                branch_name = "G-8 Main Lab (HQ)"
            
            # Normalize branch name/code lookup
            branch_key = branch_name.strip().upper()
            
            # Clean branch code generation
            branch_code = re.sub(r'[^A-Z0-9]', '', branch_key)[:10]
            if not branch_code:
                branch_code = 'GEN'
            
            # Look up branch
            branch = Branch.query.filter(
                (Branch.name.ilike(branch_name)) | (Branch.code.ilike(branch_code))
            ).first()
            
            if not branch:
                branch = Branch(
                    name=branch_name,
                    code=branch_code,
                    address=f"{branch_name} Location",
                    latitude=33.68,
                    longitude=73.03
                )
                db.session.add(branch)
                db.session.commit()
                branch_count += 1
                print(f"Created Branch: {branch_name} ({branch_code})")
            
            # Check if user already exists
            existing_user = User.query.filter(
                (User.email.ilike(email)) | (User.username == username)
            ).first()
            
            if not existing_user:
                max_emp_seq += 1
                emp_id = f"EMP-{max_emp_seq}"
                
                new_user = User(
                    employee_id=emp_id,
                    name=name,
                    username=username,
                    email=email,
                    role=role,
                    branch_id=branch.id,
                    status=True,
                    joining_date=now_pkt().date()
                )
                # Password default is username + '123'
                new_user.set_password(f"{username}123")
                db.session.add(new_user)
                user_count += 1
                print(f"Importing Employee: {name} | Email: {email} | Role: {role} | Branch: {branch.name}")
        
        db.session.commit()
        print(f"\nCompleted! Imported {user_count} new employees and created {branch_count} new branches.")

if __name__ == '__main__':
    import_staff()
