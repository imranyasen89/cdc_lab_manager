import openpyxl
from app import create_app, db
from app.models import User, StaffProfile

def import_profiles():
    app = create_app()
    with app.app_context():
        # First, ensure the tables are created!
        db.create_all()
        print("Database tables verified.")
        
        wb = openpyxl.load_workbook('CDC_lab.xlsx')
        ws = wb['cdc-staff']
        
        print(f"Reading spreadsheet: {ws.max_row - 1} entries found.")
        
        profile_count = 0
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            email = str(row[1] or '').strip()
            name = str(row[2] or '').strip()
            designation = str(row[3] or '').strip()
            
            if not email or not name:
                continue
                
            # Look up user by email
            user = User.query.filter(User.email.ilike(email)).first()
            if not user:
                # Fallback check by username (if email split username matches)
                username = email.split('@')[0].lower().replace('.', '_').replace('-', '_')
                user = User.query.filter_by(username=username).first()
                
            if not user:
                print(f"User not found for email: {email} ({name}). Skipping profile creation.")
                continue
                
            # Safe float parsing
            def parse_float(val):
                if val is None:
                    return 0.0
                try:
                    return float(val)
                except ValueError:
                    return 0.0
            
            # Check if profile already exists
            profile = StaffProfile.query.filter_by(user_id=user.id).first()
            if not profile:
                profile = StaffProfile(user_id=user.id)
                db.session.add(profile)
                
            profile.designation = designation
            profile.overall_experience = parse_float(row[6])
            profile.cdc_experience = parse_float(row[7])
            profile.qualification = str(row[8] or '').strip()
            profile.other_certificates = str(row[9] or '').strip()
            profile.sections_independent = str(row[11] or '').strip()
            profile.study_further = str(row[12] or '').strip()
            profile.research_interest = str(row[13] or '').strip()
            profile.training_requested = str(row[14] or '').strip()
            profile.job_satisfaction = str(row[15] or '').strip()
            profile.positive_aspects = str(row[16] or '').strip()
            # Suggestions column combined with any other suggestions
            sugg1 = str(row[17] or '').strip()
            sugg2 = str(row[18] or '').strip()
            profile.suggestions = f"{sugg1} | {sugg2}" if (sugg1 and sugg2) else (sugg1 or sugg2 or '')
            
            profile_count += 1
            print(f"Imported/Updated profile for: {user.name} ({user.employee_id})")
            
        db.session.commit()
        print(f"\nCompleted! Successfully imported {profile_count} staff profiles.")

if __name__ == '__main__':
    import_profiles()
